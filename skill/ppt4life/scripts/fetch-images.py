#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fetch-images.py — 用 BrowserSkill (bsk) 驱动真实浏览器搜索并下载配图，供 PPT 使用。

流程（全部在单个进程内完成，因为 bsk 的 session 不跨 shell 存活）：
  1. 启动会话   bsk session start --no-focus --json
  2. 打开图库    bsk navigate <搜索页> --session <id>
  3. 观察页面    bsk observe --session <id>      （给 agent 参考，可选）
  4. 抽取图片直链 bsk get-html --session <id>    → 正则提取图片 URL
  5. 下载图片    直链 → 本地 assets/img/（走环境代理）
  6. 结束会话    bsk session stop <id>          （finally 保证执行）

依赖：bsk CLI（BrowserSkill）+ 浏览器扩展。安装见 ../references/browser-images.md。
用法示例：
  python fetch-images.py --query "business meeting" --count 6 --out ./assets/img
  python fetch-images.py --query "科研 实验室" --source bing --count 4
  python fetch-images.py --query "product launch" --cite        # 额外生成“图片来源.pptx”
  python fetch-images.py --query "product launch" --dry-run
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request

IS_WIN = os.name == "nt"

SOURCES = {
    # 免版权图库（最合规，推荐）
    "unsplash": "https://unsplash.com/s/photos/{q}",
    "pexels": "https://www.pexels.com/search/{q}/",
    "pixabay": "https://pixabay.com/images/search/{q}/",
    # 搜索引擎图片（素材更广，注意版权）
    "bing": "https://www.bing.com/images/search?q={q}",
    "baidu": "https://image.baidu.com/search/index?tn=baiduimage&word={q}",
}

IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")


def find_bsk():
    """定位 bsk 可执行文件。"""
    cand = shutil.which("bsk")
    if cand:
        return cand
    home = os.path.expanduser("~")
    for p in (os.path.join(home, ".local", "bin", "bsk.exe" if IS_WIN else "bsk"),
              os.path.join(home, ".local", "bin", "bsk")):
        if os.path.exists(p):
            return p
    return None


def run_bsk(bsk, args, timeout):
    """执行 bsk 命令，输出重定向到临时文件（避免守护进程持有管道导致永久阻塞）。"""
    tmp = tempfile.mkdtemp(prefix="bsk_")
    out_p, err_p = os.path.join(tmp, "out.txt"), os.path.join(tmp, "err.txt")
    try:
        with open(out_p, "w", encoding="utf-8", errors="replace") as of, \
             open(err_p, "w", encoding="utf-8", errors="replace") as ef:
            p = subprocess.Popen([bsk] + args, stdout=of, stderr=ef)
            try:
                rc = p.wait(timeout=timeout)
            except subprocess.TimeoutExpired:
                if IS_WIN:
                    subprocess.run(["taskkill", "/F", "/T", "/PID", str(p.pid)],
                                   capture_output=True)
                else:
                    p.kill()
                rc = None
        out = open(out_p, encoding="utf-8", errors="replace").read()
        err = open(err_p, encoding="utf-8", errors="replace").read()
        return rc, out, err
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def extract_image_urls(html):
    """从 HTML 中提取图片 URL。图库/搜索引擎的 CDN 直链优先，其次通用图片后缀。"""
    if not html:
        return []
    patterns = [
        r'https://images\.unsplash\.com/[^\s"\'\\<>)]+',
        r'https://images\.pexels\.com/[^\s"\'\\<>)]+',
        r'https://cdn\.pixabay\.com/[^\s"\'\\<>)]+',
        r'https?://[^\s"\'\\<>)]+?\.(?:jpg|jpeg|png|webp)(?:\?[^\s"\'\\<>)]*)?',
    ]
    found, seen = [], set()
    for pat in patterns:
        for m in re.findall(pat, html, flags=re.I):
            base = m.split("?")[0]
            if base in seen:
                continue
            seen.add(base)
            found.append(m)
    return found


def safe_name(query, idx, url):
    """按 URL 路径扩展名生成安全文件名，默认 .jpg。"""
    path = url.split("?")[0].lower()
    ext = next((e for e in IMG_EXT if path.endswith(e)), ".jpg")
    slug = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "-", query).strip("-")[:30] or "img"
    return f"{slug}-{idx:02d}{ext}"


def _open(url, timeout, proxy):
    handlers = []
    if proxy:
        handlers.append(urllib.request.ProxyHandler({"http": proxy, "https": proxy}))
    else:
        handlers.append(urllib.request.ProxyHandler({}))
    opener = urllib.request.build_opener(*handlers)
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (ppt4life image fetcher)"})
    return opener.open(req, timeout=timeout)


def download(url, dest, timeout=60):
    """下载图片到本地：先走环境代理，失败再直连兜底。"""
    proxy = (os.environ.get("HTTPS_PROXY") or os.environ.get("https_proxy") or "").strip()
    last_err = None
    for p in ([proxy] if proxy else [None]) + ([None] if proxy else []):
        try:
            with _open(url, timeout, p) as r, open(dest, "wb") as f:
                shutil.copyfileobj(r, f)
            return os.path.getsize(dest)
        except Exception as e:  # 代理失败则改直连
            last_err = e
            if os.path.exists(dest):
                os.remove(dest)
    raise last_err if last_err else RuntimeError("download failed")


def write_sources_md(manifest, out_dir, query, source):
    """写人类可读的来源清单（markdown 表格），可直接贴进 PPT 备注或文档。"""
    lines = ["# 图片来源清单", "",
             f"- 关键词：{query}", f"- 图库/来源：{source}", "",
             "| # | 文件 | 来源 URL |", "|---|---|---|"]
    for i, m in enumerate(manifest, 1):
        lines.append(f"| {i} | `{m['file']}` | {m['url']} |")
    lines += ["", "> 请遵循对应图库的授权条款；若为搜索引擎图片，务必自行确认版权并在正式材料中标注出处。"]
    p = os.path.join(out_dir, "_sources.md")
    with open(p, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return p


def write_cite_slide(manifest, out_dir, query, source):
    """生成一页 16:9 的“图片来源”PPTX，可附在答辩/汇报 PPT 末尾。需 python-pptx，缺失则跳过。"""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.oxml.ns import qn
    except Exception:
        print("! 未安装 python-pptx，跳过“图片来源.pptx”生成（已生成 _sources.md）。", file=sys.stderr)
        return None

    NAVY = RGBColor(0x1F, 0x4E, 0x79); INK = RGBColor(0x26, 0x26, 0x26)
    GRAY = RGBColor(0x59, 0x59, 0x59); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    FONT = "微软雅黑"

    def setf(run, size, bold, color):
        run.font.name = FONT; run.font.size = Pt(size)
        run.font.bold = bold; run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {}); rPr.append(el)
            el.set("typeface", FONT)

    def textbox(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        r = p.add_run(); r.text = text; setf(r, size, bold, color)
        return tb

    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    textbox(s, 0.7, 0.55, 11.9, 0.7, "图片来源 · Image Credits", 24, True, NAVY)
    textbox(s, 0.7, 1.3, 11.9, 0.4,
            f"关键词：{query}    ｜    图库/来源：{source}", 12, False, GRAY)

    rows = manifest[:14]
    tbl_shape = s.shapes.add_table(len(rows) + 1, 3, Inches(0.7), Inches(1.85), Inches(11.9), Inches(4.6))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(0.8); tbl.columns[1].width = Inches(3.6); tbl.columns[2].width = Inches(7.5)
    heads = ["#", "文件", "来源 URL"]
    for c, h in enumerate(heads):
        cell = tbl.cell(0, c); cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]; r = p.add_run(); r.text = h; setf(r, 12, True, WHITE)
    for i, m in enumerate(rows, 1):
        vals = [str(i), m["file"], m["url"]]
        for c, v in enumerate(vals):
            cell = tbl.cell(i, c); cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2) if i % 2 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]; r = p.add_run()
            r.text = v if c != 2 else (v[:96] + ("…" if len(v) > 96 else ""))
            setf(r, 10, c == 0, INK if c != 2 else GRAY)

    textbox(s, 0.7, 6.65, 11.9, 0.5,
            "注：请遵循对应图库授权条款；搜索引擎图片需自行确认版权并标注出处。", 10, False, GRAY)
    out = os.path.join(out_dir, "图片来源.pptx")
    prs.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description="用 BrowserSkill 搜索并下载 PPT 配图")
    ap.add_argument("--query", required=True, help="搜索关键词（英文命中率更高）")
    ap.add_argument("--source", default="unsplash", choices=sorted(SOURCES.keys()),
                    help="图片来源（默认 unsplash 免版权图库）")
    ap.add_argument("--count", type=int, default=6, help="下载数量（默认 6）")
    ap.add_argument("--out", default="./assets/img", help="输出目录（默认 ./assets/img）")
    ap.add_argument("--timeout", type=int, default=120, help="单条 bsk 命令超时秒数")
    ap.add_argument("--cite", action="store_true",
                    help="额外生成一页 16:9“图片来源.pptx”（需 python-pptx）+ _sources.md")
    ap.add_argument("--dry-run", action="store_true", help="只打印计划，不执行")
    args = ap.parse_args()

    url = SOURCES[args.source].format(q=urllib.parse.quote(args.query))
    out_dir = os.path.abspath(args.out)
    plan = {"source": args.source, "query": args.query, "url": url,
            "count": args.count, "out": out_dir}

    if args.dry_run:
        print("[dry-run] 计划：")
        print(json.dumps(plan, ensure_ascii=False, indent=2))
        return 0

    bsk = find_bsk()
    if not bsk:
        print("✗ 未找到 bsk CLI。请先安装 BrowserSkill：", file=sys.stderr)
        print("  Windows: irm https://raw.githubusercontent.com/Tencent/BrowserSkill/main/install.ps1 | iex",
              file=sys.stderr)
        print("  详见 ../references/browser-images.md 第 0 节。", file=sys.stderr)
        return 2
    print(f"· bsk: {bsk}")

    rc, out, _ = run_bsk(bsk, ["--version"], 30)
    print(f"· bsk version: {(out or '').strip()}")

    os.makedirs(out_dir, exist_ok=True)
    sid = None
    manifest = []
    try:
        print("· 启动会话…")
        rc, out, err = run_bsk(bsk, ["session", "start", "--no-focus", "--json"], args.timeout)
        try:
            sid = json.loads(out).get("session_id")
        except Exception:
            pass
        if not sid:
            print("✗ 会话启动失败：", (out or err).strip()[:300], file=sys.stderr)
            return 3
        print(f"· session id: {sid}")

        print(f"· 打开：{url}")
        run_bsk(bsk, ["navigate", url, "--session", sid], args.timeout)

        print("· 观察页面（observe）…")
        run_bsk(bsk, ["observe", "--session", sid], args.timeout)

        print("· 抽取图片直链（get-html）…")
        _, html, _ = run_bsk(bsk, ["get-html", "--session", sid], args.timeout)
        urls = extract_image_urls(html)
        print(f"· 找到 {len(urls)} 个图片候选")
        if not urls:
            print("! 未从页面抽取到图片直链。可尝试换 --source，或在浏览器里确认页面已加载。",
                  file=sys.stderr)

        picked = urls[: max(1, args.count)]
        for i, u in enumerate(picked, 1):
            dest = os.path.join(out_dir, safe_name(args.query, i, u))
            try:
                size = download(u, dest)
                manifest.append({"file": os.path.basename(dest), "url": u,
                                 "size": size, "source": args.source})
                print(f"  ✓ {os.path.basename(dest)}  ({size//1024} KB)")
            except Exception as e:
                print(f"  ✗ 下载失败：{u[:80]}  ({e})", file=sys.stderr)

        if manifest:
            mpath = os.path.join(out_dir, "_sources.json")
            with open(mpath, "w", encoding="utf-8") as f:
                json.dump(manifest, f, ensure_ascii=False, indent=2)
            print(f"· 已写来源清单：{mpath}")
            md = write_sources_md(manifest, out_dir, args.query, args.source)
            print(f"· 已写来源说明：{md}")
            if args.cite:
                slide = write_cite_slide(manifest, out_dir, args.query, args.source)
                if slide:
                    print(f"· 已生成“图片来源”页：{slide}（可直接附在 PPT 末尾）")
        print(f"· 完成：{len(manifest)} 张图 → {out_dir}")
        return 0
    finally:
        if sid:
            print("· 结束会话…")
            run_bsk(bsk, ["session", "stop", sid], 60)


if __name__ == "__main__":
    sys.exit(main())
